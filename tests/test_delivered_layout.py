"""The rendered Word file holds together on the page.

Found paging through the Word file of a revised Chinese LCA course report:

- a five-row table started on one page and ended on the next, its caption left
  alone at the bottom of the page before it and its attribution line at the
  top of the page after it;
- a table row broke in the middle of a cell;
- every column of a parameter table got the same width, so a column of short
  numbers was as wide as a column of sentences and the table ran to five pages;
- the report title was a Heading 1, so the table of contents opened with the
  title of the document it belongs to, and the cover of a Chinese report read
  "Student Id:", "Date:" and "Keywords:".
"""
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from report_workflow.nodes.docx_render import _style_front_matter_title, _style_tables_post_render
from report_workflow.nodes.front_matter_build import _format_front_matter_markdown
from report_workflow.nodes.source_tables import render_source_table


def _png() -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(1, 1))
    ax.plot([0, 1], [0, 1])
    buffer = BytesIO()
    fig.savefig(buffer, format="png")
    plt.close(fig)
    return buffer.getvalue()


def _keep_next(paragraph) -> bool:
    return bool(paragraph.paragraph_format.keep_with_next)


class TablePaginationTests(unittest.TestCase):
    def _render(self, rows: int) -> Document:
        path = Path(tempfile.mkdtemp()) / "report.docx"
        doc = Document()
        doc.add_paragraph("表 5. 臺灣各情境的碳平衡點")
        table = doc.add_table(rows=rows + 1, cols=2)
        for index, row in enumerate(table.rows):
            row.cells[0].text = "情境" if index == 0 else f"情境 {index}"
            row.cells[1].text = "碳平衡點" if index == 0 else f"{40000 + index * 100:,}"
        doc.add_paragraph("來源：本研究計算（附錄 A24）")
        doc.save(path)
        _style_tables_post_render(str(path))
        return Document(path)

    def test_a_short_table_keeps_its_caption_rows_and_attribution_together(self):
        doc = self._render(rows=5)
        caption = doc.paragraphs[0]
        self.assertTrue(_keep_next(caption), "the caption can be left at the foot of a page")
        table = doc.tables[0]
        for row in table.rows:
            self.assertIsNotNone(row._tr.trPr.find(qn("w:cantSplit")), "a row can break inside a cell")
            for cell in row.cells:
                self.assertTrue(all(_keep_next(p) for p in cell.paragraphs), "a short table can split across pages")

    def test_a_long_table_may_break_between_rows_but_keeps_its_last_row_with_the_attribution(self):
        doc = self._render(rows=30)
        table = doc.tables[0]
        self.assertFalse(all(_keep_next(p) for p in table.rows[1].cells[0].paragraphs))
        self.assertTrue(all(_keep_next(p) for p in table.rows[-1].cells[0].paragraphs))
        self.assertTrue(all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in table.rows))

    def test_a_figure_keeps_with_its_caption(self):
        path = Path(tempfile.mkdtemp()) / "report.docx"
        doc = Document()
        doc.add_picture(BytesIO(_png()))
        doc.add_paragraph("圖 3. 臺灣基準情境下隨行駛里程累積的排放")
        doc.save(path)
        _style_tables_post_render(str(path))
        self.assertTrue(_keep_next(Document(path).paragraphs[0]))


class ColumnWidthTests(unittest.TestCase):
    def test_a_column_of_sentences_is_given_more_width_than_a_column_of_numbers(self):
        table = {
            "headers": ["參數", "數值", "說明"],
            "rows": [["電池製造排放係數（中國產 NMC622）", "80", "含電池報廢處理排放，不含回收材料抵減"],
                     ["使用年限", "15", "ICCT 歐盟研究採 20 年"]],
        }
        markdown = render_source_table(table, number=1, caption="本研究採用的參數", language="zh")
        separator = next(line for line in markdown.splitlines() if line.startswith("|") and set(line) <= set("|- :"))
        widths = [len(cell.strip()) for cell in separator.strip("|").split("|")]
        self.assertGreater(widths[2], widths[1] * 2, separator)
        self.assertGreater(widths[0], widths[1] * 2, separator)
        # pandoc only honours dash widths once a table line passes 72 columns.
        self.assertGreater(len(separator), 72)

    def test_a_header_takes_a_third_line_before_a_column_of_labels_wraps(self):
        # Capping the widest columns equally left 情境 at 3.1 cm beside a
        # header "電動車比油電混合車少排 (%)" given as much: three labels wrapped,
        # the table outgrew the space left on its page and moved, and a figure
        # followed it — two half-empty pages for one header line saved.
        from report_workflow.nodes import source_tables

        table = {
            "headers": ["情境", "汽油車 (t CO2e)", "油電混合車 (t CO2e)", "電動車 (t CO2e)",
                        "電動車比汽油車少排 (%)", "電動車比油電混合車少排 (%)"],
            "rows": [["A 只算燃燒", "36.3", "29.8", "28.3", "22.0", "4.8"],
                     ["歐盟潔淨電網（參考）", "44.2", "35.8", "14.3", "67.6", "60.1"]],
        }
        markdown = render_source_table(table, number=4, caption="臺灣各情境 15 年全生命週期排放", language="zh")
        separator = next(line for line in markdown.splitlines() if line.startswith("|") and set(line) <= set("|- :"))
        dashes = [len(cell.strip()) for cell in separator.strip("|").split("|")]
        label_cm = dashes[0] / sum(dashes) * source_tables.TEXT_WIDTH_CM
        needed_cm = source_tables._display_width("歐盟潔淨電網（參考）") * source_tables.DISPLAY_COLUMN_CM + source_tables.CELL_PADDING_CM
        self.assertGreaterEqual(label_cm, needed_cm - 0.05, separator)

    def test_widths_are_estimated_at_the_type_size_the_template_prints(self):
        # Widths were estimated for 10.5 pt text and the template prints tables
        # at 12 pt, so a unit header came out as "電力係數 (g" / "CO2e/kWh" / ")".
        import re
        import zipfile

        from report_workflow.nodes import source_tables

        template = Path(source_tables.__file__).parent.parent / "templates" / "reference.docx"
        styles = zipfile.ZipFile(template).read("word/styles.xml").decode("utf-8")
        defaults = re.search(r"<w:rPrDefault>.*?</w:rPrDefault>", styles, re.S).group(0)
        points = int(re.search(r'<w:sz w:val="(\d+)"', defaults).group(1)) / 2
        half_character_cm = points / 2 / 72 * 2.54
        self.assertGreaterEqual(source_tables.DISPLAY_COLUMN_CM, round(half_character_cm, 3))


class FrontMatterTests(unittest.TestCase):
    def test_a_chinese_cover_is_labelled_in_chinese(self):
        markdown = _format_front_matter_markdown({
            "title": "電動車與燃油車全生命週期二氧化碳排放評估",
            "author_block": "蔡知均",
            "template_fields": {"student_id": "E14122313", "date": "2026年9月"},
            "keywords": ["生命週期評估", "電動車"],
        })
        self.assertIn("**學號：** E14122313", markdown)
        self.assertIn("**日期：** 2026年9月", markdown)
        self.assertIn("**關鍵詞：** 生命週期評估、電動車", markdown)
        self.assertNotIn("Student Id", markdown)

    def test_the_field_report_looks_for_the_keywords_as_the_cover_prints_them(self):
        # Joined with 、 on a Chinese cover, the keywords were then searched for
        # comma-joined, and the delivery summary reported them missing.
        from report_workflow.artifact_packaging.template_reports import _front_matter_field_values

        fields = _front_matter_field_values({"title": "電動車評估", "keywords": ["生命週期評估", "電動車"]})
        self.assertEqual(fields["keywords"]["value"], "生命週期評估、電動車")

    def test_an_english_cover_is_unchanged(self):
        markdown = _format_front_matter_markdown({"title": "Capstone Lab Report", "template_fields": {"student_id": "S1"}})
        self.assertIn("**Student Id:** S1", markdown)

    def test_a_long_title_breaks_after_its_colon_not_inside_a_word(self):
        # The cover printed "…全生命週期二氧化" / "碳排放評估".
        for title, printed in (
            ("從製造、使用到回收：電動車與燃油車全生命週期二氧化碳排放評估",
             "從製造、使用到回收：\n電動車與燃油車全生命週期二氧化碳排放評估"),
            ("電動車評估：臺灣情境", "電動車評估：臺灣情境"),
        ):
            with self.subTest(title=title):
                path = Path(tempfile.mkdtemp()) / "report.docx"
                doc = Document()
                doc.add_heading(title, level=1)
                doc.save(path)
                self.assertTrue(_style_front_matter_title(str(path), title))
                self.assertEqual(Document(path).paragraphs[0].text, printed)

    def test_the_field_report_still_finds_a_title_printed_on_two_lines(self):
        from report_workflow.artifact_packaging.template_reports import _rendered_matches

        rendered = "從製造、使用到回收：\n電動車與燃油車全生命週期二氧化碳排放評估\n蔡知均"
        self.assertEqual(_rendered_matches(rendered, "從製造、使用到回收：電動車與燃油車全生命週期二氧化碳排放評估"), 1)
        self.assertEqual(_rendered_matches("Capstone  Lab\nReport", "Capstone Lab Report"), 1)

    def test_the_title_is_styled_as_a_title_not_a_heading(self):
        path = Path(tempfile.mkdtemp()) / "report.docx"
        doc = Document()
        doc.add_heading("電動車與燃油車全生命週期二氧化碳排放評估", level=1)
        doc.add_heading("1. 緒論", level=1)
        doc.save(path)
        _style_front_matter_title(str(path), "電動車與燃油車全生命週期二氧化碳排放評估")
        styled = Document(path)
        self.assertEqual(styled.paragraphs[0].style.name, "Title")
        self.assertEqual(styled.paragraphs[1].style.name, "Heading 1")


if __name__ == "__main__":
    unittest.main()
