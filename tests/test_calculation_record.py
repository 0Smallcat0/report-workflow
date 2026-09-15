"""The report explains its calculations; the full record travels beside it.

Second revision of a Chinese LCA course report. The document closed on seven
pages of generated entries — g_cfp and bev_body for variables, "用Ember 0、0、1、
0、0" switches, the same eleven parameters and a four-line formula repeated
under every table — and the reviewer asked for an appendix a reader can use:
the formulas, the parameters and one worked example, with the complete record
saved as a separate attachment. So the generated record is rendered as its own
document, the report keeps the appendix the author writes, references stay in
front of that appendix, and the appendix heading is not numbered as a chapter.
"""
import json
import unittest
from unittest.mock import patch

from docx import Document

from report_workflow.derived_evidence import build_requested_units
from report_workflow.nodes.docx_render import _insert_before_appendices, run_docx_render
from report_workflow.nodes.heading_contract_check import normalize_heading_contract
from report_workflow.state import WORKFLOW_RUNS_DIR, ReportState

from tests.test_formula_derivations import REGISTRY, _cell

CREATED_AT = "2026-09-16T00:00:00+00:00"


class CalculationRecordTests(unittest.TestCase):
    @patch("report_workflow.nodes.docx_render._validate_docx", return_value=[])
    @patch("report_workflow.nodes.docx_render._style_tables_post_render", return_value=None)
    @patch("report_workflow.nodes.docx_render._repair_missing_figures", return_value=0)
    def test_the_record_is_its_own_document_not_the_end_of_the_report(self, _repair, _style, _validate):
        state = ReportState.new("report", [], "out")
        state.qa["qa_decision"] = "pass"
        state.spec["report_profile"] = "custom"
        run_dir = WORKFLOW_RUNS_DIR / state.job_id
        run_dir.mkdir(parents=True, exist_ok=True)
        draft = run_dir / "publication.md"
        draft.write_text(
            "# 結果\n\n在本研究的臺灣基準參數下，電動車的製造與回收排放比汽油車多 3.572 t CO2e"
            "（本研究計算，計算附件 C1），差距幾乎全部來自電池。\n",
            encoding="utf-8",
        )
        state.drafts["publication_style_draft"] = str(draft)
        record = run_dir / "publication_source_list.md"
        record.write_text(
            "# Calculation Record\n\nintro\n\n- **C1**　電動車比汽油車多出的製造與回收排放為 3.572 t CO2e（本研究計算）。"
            "計算式：bev_body + battery - ice_body。\n",
            encoding="utf-8",
        )
        state.citations["publication_source_list_path"] = str(record)

        rendered: list[str] = []

        def fake_pandoc(md_path, output_path, reference_doc=None, warnings_out=None):
            rendered.append(open(md_path, encoding="utf-8").read())
            doc = Document()
            doc.add_heading("結果", level=1)
            doc.add_paragraph("rendered")
            doc.save(output_path)
            return True

        with patch("report_workflow.nodes.docx_render._render_via_pandoc", side_effect=fake_pandoc):
            result = run_docx_render(state)

        report_input = (run_dir / "pandoc_input.md").read_text(encoding="utf-8")
        self.assertNotIn("bev_body + battery", report_input)
        self.assertNotIn("Calculation Record", report_input)
        record_docx = result.output.get("calculation_record_docx_path")
        self.assertTrue(record_docx, result.output)
        record_md = next(text for text in rendered if "bev_body" in text)
        self.assertIn("# 計算附件　完整運算紀錄", record_md)
        self.assertIn("**C1**", record_md)


class AppendixPlacementTests(unittest.TestCase):
    def test_references_come_before_an_appendix_the_author_wrote(self):
        body = "# 5. 建議事項\n\n結論。\n\n# 附錄\n\n## A.1 計算式\n\n式 A1。\n"
        placed = _insert_before_appendices(body, "# 參考文獻\n\n- 文獻甲\n")
        self.assertLess(placed.index("# 參考文獻"), placed.index("# 附錄"))
        self.assertLess(placed.index("結論。"), placed.index("# 參考文獻"))

    def test_without_an_appendix_references_close_the_document(self):
        placed = _insert_before_appendices("# 5. 建議事項\n\n結論。\n", "# 參考文獻\n\n- 文獻甲\n")
        self.assertTrue(placed.rstrip().endswith("- 文獻甲"), placed)

    def test_an_appendix_is_not_numbered_as_a_chapter(self):
        blueprint = {
            "section_order": ["introduction", "references", "appendix"],
            "sections": {
                "introduction": {"title": "Introduction", "title_zh": "緒論", "required": True},
                "references": {"title": "References", "title_zh": "參考文獻", "required": True},
                "appendix": {"title": "Appendix", "title_zh": "附錄", "required": False},
            },
        }
        normalized, _issues = normalize_heading_contract(
            "# 緒論\n\n本研究以生命週期評估比較電動車與燃油車在製造、使用與回收階段的排放。\n\n"
            "# 參考文獻\n\n# 附錄\n\n本附錄說明計算式、參數與基準情境的代入範例。\n",
            blueprint,
        )
        self.assertIn("# 1. 緒論", normalized)
        self.assertIn("# 附錄", normalized)
        self.assertNotIn("# 2. 附錄", normalized)


class ExplainedInTests(unittest.TestCase):
    def test_a_derivation_keeps_where_the_report_explains_it(self):
        units, problems = build_requested_units(
            [{"id": "battery_t", "op": "formula", "expr": "kwh * factor / 1000", "unit": "t CO2e", "decimals": 3,
              "inputs": {"kwh": _cell("電池容量"), "factor": _cell("電池製造係數")}, "explained_in": "附錄 A.3"}],
            REGISTRY, CREATED_AT, zh=True,
        )
        self.assertEqual(problems, [])
        self.assertEqual(units[0]["derivation"]["explained_in"], "附錄 A.3")

    def test_a_location_that_is_not_text_is_refused(self):
        _units, problems = build_requested_units(
            [{"id": "battery_t", "op": "formula", "expr": "kwh * factor / 1000", "unit": "t CO2e",
              "inputs": {"kwh": _cell("電池容量"), "factor": _cell("電池製造係數")}, "explained_in": 3}],
            REGISTRY, CREATED_AT, zh=True,
        )
        self.assertEqual(len(problems), 1, problems)
        self.assertIn("explained_in", json.dumps(problems, ensure_ascii=False))


if __name__ == "__main__":
    unittest.main()
