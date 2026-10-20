"""check-references reads a draft written anywhere and checks every reference in it.

The registry comparison itself is pinned in test_reference_match; these tests
pin what feeds it -- the reference list pulled out of a .md, .bib or .docx --
and what is done with its verdicts: the command's output and exit code, and
REFERENCE_VERIFY's report inside a workflow run. No test touches the network.
"""
import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

from report_workflow import reference_match as rm
from report_workflow.cli import main as cli_main
from report_workflow.errors import QAHardBlockError
from report_workflow.nodes.reference_reality_check import run_reference_reality_check
from report_workflow.nodes.reference_verify import run_reference_verify
from report_workflow.reference_list import extract_references
from report_workflow.state import ReportState

MARKDOWN = """# Findings

Static single assignment form [1] is assumed throughout.

## References

1. Cytron, R., Ferrante, J., Rosen, B. K., Wegman, M. N., & Zadeck, F. K. (1991).
   Efficiently computing static single assignment form and the control
   dependence graph. ACM TOPLAS, 13(4), 451-490. https://doi.org/10.1145/115372.115320
2. Touvron, H., Lavril, T., et al. (2023). LLaMA: Open and efficient foundation
   language models. arXiv:2302.13971

## Appendix A

Not a reference, and long enough to be mistaken for one.
"""

BIBTEX = """@article{cytron1991,
  author = {Cytron, Ron and Ferrante, Jeanne},
  title = {Efficiently Computing Static Single Assignment Form},
  journal = {ACM TOPLAS},
  year = {1991},
  doi = {10.1145/115372.115320}
}

@misc{touvron2023,
  author = {Touvron, Hugo and others},
  title = {{LLaMA}: Open and Efficient Foundation Language Models},
  year = {2023},
  eprint = {2302.13971},
  archivePrefix = {arXiv}
}
"""


def verdicts(table):
    """A stand-in for check_reference that answers by a substring of the reference."""

    def check(reference, **_kwargs):
        for needle, status in table.items():
            if needle in reference:
                return rm.ReferenceCheck(status, f"{status} for {needle}", identifier="doi:10.1/x")
        return rm.ReferenceCheck(rm.NEEDS_REVIEW, "nothing to compare")

    return check


class ExtractTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_markdown_reads_only_the_reference_section_and_joins_wrapped_lines(self):
        path = self.tmp / "draft.md"
        path.write_text(MARKDOWN, encoding="utf-8")
        refs = extract_references(path)
        self.assertEqual(len(refs), 2)
        self.assertIn("single assignment form and the control dependence graph.", refs[0])
        self.assertTrue(refs[0].startswith("Cytron, R."))
        self.assertTrue(refs[1].endswith("arXiv:2302.13971"))

    def test_bibtex_entries_read_as_references_with_their_identifiers(self):
        path = self.tmp / "refs.bib"
        path.write_text(BIBTEX, encoding="utf-8")
        refs = extract_references(path)
        self.assertEqual(len(refs), 2)
        self.assertTrue(refs[0].startswith("Ron Cytron, Jeanne Ferrante (1991)."))
        self.assertIn("doi:10.1145/115372.115320", refs[0])
        self.assertIn("LLaMA: Open and Efficient", refs[1])
        self.assertIn("arXiv:2302.13971", refs[1])
        # The names come out the way a reference writes them.
        self.assertEqual(rm.cited_surnames(refs[0].split("(")[0]), ["Cytron", "Ferrante"])

    def test_docx_takes_one_paragraph_per_reference_under_the_heading(self):
        document = Document()
        document.add_heading("Results", level=1)
        document.add_paragraph("Body text that cites prior work at some length here.")
        document.add_heading("References", level=1)
        document.add_paragraph("Cytron, R. (1991). Efficiently computing static single assignment form.")
        document.add_paragraph("Touvron, H. (2023). LLaMA: Open and efficient foundation language models.")
        document.add_heading("Appendix", level=1)
        document.add_paragraph("Appendix text that is not a reference at all, but is long.")
        path = self.tmp / "draft.docx"
        document.save(str(path))
        refs = extract_references(path)
        self.assertEqual([r[:7] for r in refs], ["Cytron,", "Touvron"])

    def test_a_file_with_no_heading_is_the_list_itself(self):
        path = self.tmp / "list.txt"
        path.write_text(
            "Cytron, R. (1991). Efficiently computing static single assignment form.\n\n"
            "Touvron, H. (2023). LLaMA: Open and efficient foundation language models.\n",
            encoding="utf-8",
        )
        self.assertEqual(len(extract_references(path)), 2)


class CommandTest(unittest.TestCase):
    def run_cli(self, *argv):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = cli_main(["check-references", *argv])
        return code, out.getvalue(), err.getvalue()

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.draft = self.tmp / "draft.md"
        self.draft.write_text(MARKDOWN, encoding="utf-8")

    def test_a_refuted_reference_fails_the_command_and_is_named(self):
        table = {"Cytron": rm.VERIFIED, "Touvron": rm.MISMATCH}
        with mock.patch.object(rm, "check_reference", verdicts(table)):
            code, out, _ = self.run_cli(str(self.draft), "--json", str(self.tmp / "out.json"))
        self.assertEqual(code, 2)
        self.assertIn("[2] MISMATCH: mismatch for Touvron", out)
        self.assertIn("1 mismatch, 1 verified (2 references)", out)
        written = json.loads((self.tmp / "out.json").read_text(encoding="utf-8"))
        self.assertEqual([item["status"] for item in written], ["verified", "mismatch"])
        self.assertEqual(written[1]["number"], 2)

    def test_references_left_for_review_do_not_fail_it(self):
        table = {"Cytron": rm.VERIFIED, "Touvron": rm.NEEDS_REVIEW}
        with mock.patch.object(rm, "check_reference", verdicts(table)):
            code, out, _ = self.run_cli(str(self.draft))
        self.assertEqual(code, 0)
        self.assertIn("1 needs_review", out)

    def test_no_search_is_passed_through(self):
        seen = []

        def check(reference, *, search=True, **_kwargs):
            seen.append(search)
            return rm.ReferenceCheck(rm.VERIFIED, "ok")

        with mock.patch.object(rm, "check_reference", check):
            self.run_cli(str(self.draft), "--no-search")
        self.assertEqual(seen, [False, False])

    def test_a_missing_file_is_an_error_not_a_crash(self):
        code, _, err = self.run_cli(str(self.tmp / "nope.md"))
        self.assertEqual(code, 1)
        self.assertIn("No such file", err)


class WorkflowNodeTest(unittest.TestCase):
    """REFERENCE_VERIFY records each verdict, and the reality check routes it."""

    def make_state(self, profile):
        tmp = Path(tempfile.mkdtemp())
        state = ReportState.new("report", [], str(tmp / "out"))
        state.spec["report_profile"] = profile
        ref_list = tmp / "publication_reference_list.md"
        ref_list.write_text(
            "# References\n\n"
            "- Cytron, R. (1991). Efficiently computing SSA form. ACM TOPLAS. doi:10.1145/115372.115320\n"
            "- Nguyen, T. (2021). Adaptive caching for serverless inference. doi:10.1145/115372.115320\n"
            "- Kong, A. (1992). A note on importance sampling. Journal of Statistics, 3(1), 1-9. doi:10.1/kong\n",
            encoding="utf-8",
        )
        state.citations["publication_reference_list_path"] = str(ref_list)
        return state

    def test_a_mismatch_hard_blocks_an_academic_paper_and_review_is_routed(self):
        state = self.make_state("academic_paper")
        table = {"Cytron": rm.VERIFIED, "Nguyen": rm.MISMATCH, "Kong": rm.NEEDS_REVIEW}
        with mock.patch("report_workflow.nodes.reference_verify.check_reference", verdicts(table)):
            with self.assertRaises(QAHardBlockError) as raised:
                run_reference_verify(state)
        self.assertIn("mismatch for Nguyen", str(raised.exception))
        report = json.loads(Path(state.runtime["reference_verify_report_path"]).read_text(encoding="utf-8"))
        self.assertEqual((report["verified"], report["failed"], report["needs_review"]), (1, 1, 1))

    def test_review_reaches_the_reality_check_as_human_review(self):
        state = self.make_state("business_report")
        table = {"Cytron": rm.VERIFIED, "Nguyen": rm.VERIFIED, "Kong": rm.NEEDS_REVIEW}
        with mock.patch("report_workflow.nodes.reference_verify.check_reference", verdicts(table)):
            run_reference_verify(state)
        run_reference_reality_check(state)
        reality = json.loads(Path(state.runtime["reference_reality_report_path"]).read_text(encoding="utf-8"))
        self.assertEqual(reality["status"], "needs_human_review")
        self.assertEqual([item["reason"] for item in reality["needs_human_review"]], ["needs_review for Kong"])
        categories = sorted(item["category"] for item in reality["classified_references"])
        self.assertEqual(categories, ["needs_human_review", "verified_external_identifier", "verified_external_identifier"])

    def test_search_follows_web_research_unless_set_explicitly(self):
        seen = []

        def check(reference, *, search=True, **_kwargs):
            seen.append(search)
            return rm.ReferenceCheck(rm.VERIFIED, "ok", identifier="doi:10.1/x")

        for flags, expected in (({}, False), ({"enable_research": True}, True),
                                ({"enable_research": True, "reference_search": False}, False)):
            state = self.make_state("business_report")
            state.flags.update(flags)
            seen.clear()
            with mock.patch("report_workflow.nodes.reference_verify.check_reference", check):
                run_reference_verify(state)
            self.assertEqual(set(seen), {expected}, flags)


if __name__ == "__main__":
    unittest.main()
